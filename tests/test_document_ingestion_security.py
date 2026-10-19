"""Deterministic, local tests for bounded document ingestion."""

from __future__ import annotations

import io
import struct
import zipfile

import pytest
from langchain_core.embeddings import Embeddings

from module5.document_upload import (
    DEFAULT_MAX_DOCX_COMPRESSION_RATIO,
    DEFAULT_MAX_DOCX_ENTRIES,
    DEFAULT_MAX_DOCX_MEMBER_BYTES,
    DEFAULT_MAX_DOCX_UNCOMPRESSED_BYTES,
    DEFAULT_MAX_EXTRACTED_TEXT_BYTES_PER_DOCUMENT,
    DEFAULT_MAX_EXTRACTED_TEXT_BYTES_PER_UPLOAD,
    DEFAULT_MAX_EXTRACTED_TEXT_CHARS_PER_DOCUMENT,
    DEFAULT_MAX_EXTRACTED_TEXT_CHARS_PER_UPLOAD,
    DEFAULT_MAX_PDF_PAGES,
    DocumentUploader,
    UploadValidationError,
    _safe_basename,
)


def _make_pdf(texts: list[str]) -> bytes:
    """Create a small searchable PDF entirely in memory."""
    import pymupdf

    document = pymupdf.open()
    try:
        for text in texts:
            page = document.new_page()
            page.insert_text((72, 72), text)
        return document.tobytes()
    finally:
        document.close()


def _make_docx(paragraphs: list[str]) -> bytes:
    """Create a minimal DOCX in memory without filesystem extraction."""
    from docx import Document as DocxDocument

    document = DocxDocument()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _make_zip(entries: list[tuple[str, bytes]]) -> bytes:
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, value in entries:
            archive.writestr(name, value)
    return stream.getvalue()


def _rewrite_first_central_sizes(
    payload: bytes, *, file_size: int | None = None, compressed_size: int | None = None
) -> bytes:
    """Change only central-directory metadata; the validator must reject before reads."""
    mutated = bytearray(payload)
    offset = mutated.find(b"PK\x01\x02")
    assert offset >= 0, "fixture must contain a central-directory entry"
    if compressed_size is not None:
        struct.pack_into("<I", mutated, offset + 20, compressed_size)
    if file_size is not None:
        struct.pack_into("<I", mutated, offset + 24, file_size)
    return bytes(mutated)


class DeterministicEmbeddings(Embeddings):
    """Local embeddings for transactional indexing tests; no model/API calls."""

    def embed_documents(self, texts):
        return [self.embed_query(text) for text in texts]

    def embed_query(self, text):
        return [float(len(text)), float(sum(text.encode("utf-8")) % 101), 1.0]


def test_default_8_mib_upload_2_mib_document_and_200_page_limits_are_explicit():
    assert DEFAULT_MAX_EXTRACTED_TEXT_BYTES_PER_UPLOAD == 8 * 1024 * 1024
    assert DEFAULT_MAX_EXTRACTED_TEXT_CHARS_PER_UPLOAD == 2_000_000
    assert DEFAULT_MAX_EXTRACTED_TEXT_BYTES_PER_DOCUMENT == 2 * 1024 * 1024
    assert DEFAULT_MAX_EXTRACTED_TEXT_CHARS_PER_DOCUMENT == 500_000
    assert DEFAULT_MAX_PDF_PAGES == 200


def test_default_docx_limits_are_1024_entries_50_mib_10_mib_and_100_to_1():
    assert DEFAULT_MAX_DOCX_ENTRIES == 1_024
    assert DEFAULT_MAX_DOCX_UNCOMPRESSED_BYTES == 50 * 1024 * 1024
    assert DEFAULT_MAX_DOCX_MEMBER_BYTES == 10 * 1024 * 1024
    assert DEFAULT_MAX_DOCX_COMPRESSION_RATIO == 100.0


def test_safe_basename_rejects_dot_components_and_windows_traversal():
    assert _safe_basename("../../etc/passwd") == "passwd"
    assert _safe_basename(r"C:\Users\bob\report.pdf") == "report.pdf"
    assert _safe_basename("..") == "uploaded"
    assert _safe_basename("/") == "uploaded"


def test_text_limits_bound_raw_bytes_before_decode_and_parsed_document_chars():
    raw_limited = DocumentUploader(max_extracted_text_bytes_per_upload=16)
    with pytest.raises(UploadValidationError, match="per-upload byte limit"):
        raw_limited.parse_file(b"a" * 17, "large.txt")

    document_limited = DocumentUploader(
        max_extracted_text_bytes_per_upload=1_000,
        max_extracted_text_chars_per_upload=1_000,
        max_extracted_text_chars_per_document=12,
    )
    with pytest.raises(UploadValidationError, match="per-document character limit"):
        document_limited.parse_file(b"0123456789ABC", "large-paragraph.txt")


def test_text_limits_bound_total_extracted_utf8_bytes_and_chars_per_upload():
    uploader = DocumentUploader(
        max_extracted_text_bytes_per_upload=20,
        max_extracted_text_chars_per_upload=20,
        max_extracted_text_bytes_per_document=20,
        max_extracted_text_chars_per_document=20,
    )
    with pytest.raises(UploadValidationError, match="per-upload"):
        uploader.parse_file(b"abcdefghijk\n\nlmnopqrstuv", "two-paragraphs.txt")

    unicode_uploader = DocumentUploader(
        max_extracted_text_bytes_per_upload=1_000,
        max_extracted_text_chars_per_upload=1_000,
        max_extracted_text_bytes_per_document=12,
        max_extracted_text_chars_per_document=100,
    )
    with pytest.raises(UploadValidationError, match="per-document byte limit"):
        unicode_uploader.parse_file(("é" * 11).encode("utf-8"), "unicode.txt")


def test_text_parser_keeps_normal_paragraphs_without_materializing_split_list():
    uploader = DocumentUploader(
        max_extracted_text_bytes_per_upload=200,
        max_extracted_text_chars_per_upload=200,
        max_extracted_text_bytes_per_document=100,
        max_extracted_text_chars_per_document=100,
    )
    docs = uploader.parse_file(
        b"first bounded paragraph\n\nsecond bounded paragraph", "notes.txt"
    )
    assert [doc.page_content for doc in docs] == [
        "first bounded paragraph",
        "second bounded paragraph",
    ]


def test_malformed_pdf_is_rejected_and_does_not_publish_a_filename():
    uploader = DocumentUploader()
    with pytest.raises(UploadValidationError, match="PDF"):
        uploader.parse_file(b"%PDF-1.7\nnot a valid PDF", "bad.pdf")
    assert uploader.filename == ""


def test_pdf_page_count_limit_is_checked_before_text_extraction():
    uploader = DocumentUploader(max_pdf_pages=1)
    with pytest.raises(UploadValidationError, match="page count"):
        uploader.parse_file(_make_pdf(["first page", "second page"]), "many-pages.pdf")


def test_pdf_extracted_text_is_bounded_per_page_and_per_upload():
    page_limited = DocumentUploader(
        max_extracted_text_bytes_per_upload=1_000,
        max_extracted_text_chars_per_upload=1_000,
        max_extracted_text_bytes_per_document=20,
        max_extracted_text_chars_per_document=20,
    )
    with pytest.raises(UploadValidationError, match="per-document"):
        page_limited.parse_file(_make_pdf(["a page with more than twenty chars"]), "page.pdf")

    upload_limited = DocumentUploader(
        max_extracted_text_bytes_per_upload=20,
        max_extracted_text_chars_per_upload=20,
        max_extracted_text_bytes_per_document=100,
        max_extracted_text_chars_per_document=100,
    )
    with pytest.raises(UploadValidationError, match="per-upload"):
        upload_limited.parse_file(_make_pdf(["first page text", "second page text"]), "pages.pdf")


def test_pdf_parser_closes_parser_handle_when_page_extraction_fails(monkeypatch):
    import pymupdf

    class FailingPage:
        def get_text(self, mode):
            raise RuntimeError("synthetic page failure")

    class TrackedDocument:
        is_encrypted = False

        def __init__(self):
            self.closed = False

        def __len__(self):
            return 1

        def __iter__(self):
            return iter([FailingPage()])

        def close(self):
            self.closed = True

    tracked = TrackedDocument()
    monkeypatch.setattr(pymupdf, "open", lambda **kwargs: tracked)
    uploader = DocumentUploader()
    with pytest.raises(UploadValidationError, match="PDF"):
        uploader.parse_file(b"%PDF-1.7\nsynthetic", "broken.pdf")
    assert tracked.closed


@pytest.mark.parametrize("member_name", ["../evil.txt", "/etc/passwd", r"C:\evil.txt", r"a\..\evil.txt"])
def test_docx_rejects_absolute_and_traversal_zip_member_names(member_name):
    uploader = DocumentUploader()
    with pytest.raises(UploadValidationError, match="unsafe path"):
        uploader.parse_file(_make_zip([(member_name, b"payload")]), "unsafe.docx")


def test_docx_rejects_too_many_zip_entries_before_python_docx(monkeypatch):
    import docx

    called = False

    def should_not_parse(*args, **kwargs):
        nonlocal called
        called = True
        raise AssertionError("python-docx must not receive a rejected archive")

    monkeypatch.setattr(docx, "Document", should_not_parse)
    payload = _make_zip([("one", b"1"), ("two", b"2")])
    with pytest.raises(UploadValidationError, match="too many ZIP entries"):
        DocumentUploader(max_docx_entries=1).parse_file(payload, "many.docx")
    assert not called


def test_docx_rejects_oversized_member_and_total_expansion_metadata():
    member_payload = _rewrite_first_central_sizes(
        _make_zip([("word/document.xml", b"x")]), file_size=101
    )
    with pytest.raises(UploadValidationError, match="member .* too large"):
        DocumentUploader(max_docx_member_bytes=100).parse_file(member_payload, "member.docx")

    total_payload = _make_zip([("one", b"123456"), ("two", b"123456")])
    with pytest.raises(UploadValidationError, match="total uncompressed"):
        DocumentUploader(
            max_docx_member_bytes=10,
            max_docx_uncompressed_bytes=10,
        ).parse_file(total_payload, "total.docx")


def test_docx_rejects_suspicious_compression_ratio_from_central_metadata():
    payload = _rewrite_first_central_sizes(
        _make_zip([("word/document.xml", b"x")]), file_size=1_001, compressed_size=1
    )
    with pytest.raises(UploadValidationError, match="compression ratio"):
        DocumentUploader(max_docx_compression_ratio=10).parse_file(payload, "ratio.docx")


def test_malformed_docx_zip_is_rejected_before_python_docx():
    with pytest.raises(UploadValidationError, match="DOCX ZIP"):
        DocumentUploader().parse_file(b"PK\x03\x04not-a-zip", "malformed.docx")


def test_docx_extracted_text_limit_applies_after_safe_archive_inspection():
    payload = _make_docx(["this paragraph is longer than the configured bound"])
    with pytest.raises(UploadValidationError, match="per-document character limit"):
        DocumentUploader(max_extracted_text_chars_per_document=12).parse_file(
            payload, "large.docx"
        )


def test_default_docx_limits_allow_a_small_valid_in_memory_document():
    docs = DocumentUploader().parse_file(
        _make_docx(["A normal DOCX paragraph remains readable after ZIP inspection."]),
        "normal.docx",
    )
    assert docs[0].page_content.startswith("A normal DOCX paragraph")


def test_failed_security_parse_preserves_previous_valid_upload_and_collection(monkeypatch):
    import module5.document_upload as document_upload

    live_collections = []

    class TrackingVectorStore:
        def __init__(self, collection_name, embedding_function):
            self.name = collection_name
            self.deleted = False
            live_collections.append(self)

        def add_documents(self, documents):
            self.documents = list(documents)

        def delete_collection(self):
            self.deleted = True
            live_collections.remove(self)

    monkeypatch.setattr(document_upload, "Chroma", TrackingVectorStore)
    uploader = DocumentUploader(
        DeterministicEmbeddings(),
        max_extracted_text_chars_per_document=32,
    )
    uploader.index_documents(b"Previous valid upload text.", "previous.txt")
    previous_store = uploader.vectorstore
    previous_chunks = uploader.chunks

    try:
        with pytest.raises(UploadValidationError, match="per-document"):
            uploader.index_documents(
                b"This replacement is deliberately too long.",
                "replacement.txt",
                # The parser limit is configured on the uploader, so this
                # remains a parse failure before a replacement collection exists.
            )
        assert uploader.vectorstore is previous_store
        assert uploader.chunks == previous_chunks
        assert uploader.filename == "previous.txt"
        assert live_collections == [previous_store]
    finally:
        uploader.clear()
