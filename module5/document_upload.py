# module5/document_upload.py
"""
Document Upload — Let users upload their own PDF/TXT/DOCX/CSV/MD and query it.
Creates a temporary in-memory vector store from the parsed content.

Security guards:
  • file size capped by configs.config.MAX_UPLOAD_BYTES
  • magic-byte sniffing + extension allow-list (defence in depth)
  • path-traversal-safe filename storage
  • bounded extracted text, PDF pages, and DOCX ZIP expansion
"""

from __future__ import annotations

import logging
import io
import math
import os
import sys
import zipfile
from copy import copy
from pathlib import PurePosixPath, PureWindowsPath
from typing import List, Optional
from uuid import uuid4

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from langchain_community.vectorstores import Chroma
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from configs.config import CHUNK_OVERLAP, CHUNK_SIZE, MAX_UPLOAD_BYTES

logger = logging.getLogger(__name__)


# Parser limits intentionally live here rather than in the global application
# configuration: they are upload-parser policy, and tests/callers can override
# them on one DocumentUploader without changing unrelated application settings.
# The defaults cap extracted text at 8 MiB/2 million characters per upload and
# 2 MiB/500,000 characters per emitted Document.  The latter prevents one PDF
# page, paragraph, or table row from becoming an unexpectedly large chunk before
# the normal text splitter runs.
DEFAULT_MAX_EXTRACTED_TEXT_BYTES_PER_UPLOAD = 8 * 1024 * 1024
DEFAULT_MAX_EXTRACTED_TEXT_CHARS_PER_UPLOAD = 2_000_000
DEFAULT_MAX_EXTRACTED_TEXT_BYTES_PER_DOCUMENT = 2 * 1024 * 1024
DEFAULT_MAX_EXTRACTED_TEXT_CHARS_PER_DOCUMENT = 500_000
DEFAULT_MAX_PDF_PAGES = 200
DEFAULT_MAX_PARSED_DOCUMENTS = 10_000

# DOCX is a ZIP container.  These bounds are checked from central-directory
# metadata before python-docx sees the archive, so no member is extracted to
# disk.  A 50 MiB expanded archive is enough for ordinary text documents while
# remaining below the upload/parser memory budget; a 100:1 ratio rejects the
# highly-compressible metadata commonly used by ZIP bombs without rejecting
# normal DOCX XML/media members.
DEFAULT_MAX_DOCX_ENTRIES = 1_024
DEFAULT_MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
DEFAULT_MAX_DOCX_MEMBER_BYTES = 10 * 1024 * 1024
DEFAULT_MAX_DOCX_COMPRESSION_RATIO = 100.0


# Magic-byte signatures for the file types we accept.
_MAGIC_BYTES = {
    "pdf":  (b"%PDF",),
    "docx": (b"PK\x03\x04",),  # DOCX is a ZIP container
}


def _safe_basename(name: str) -> str:
    """
    Reduce an uploaded filename to a safe basename.
    Strips any directory components and drops control characters.
    """
    base = PurePosixPath(name.replace("\\", "/")).name
    base = "".join(c for c in base if 32 <= ord(c) < 127)
    return base if base and base not in {".", ".."} else "uploaded"


class UploadValidationError(ValueError):
    """Raised when an uploaded file fails security validation."""


def _trimmed_span(text: str, start: int = 0, end: Optional[int] = None) -> tuple[int, int]:
    """Return non-whitespace bounds without copying the text."""
    if end is None:
        end = len(text)
    while start < end and text[start].isspace():
        start += 1
    while end > start and text[end - 1].isspace():
        end -= 1
    return start, end


class _TextBudget:
    """Account for bounded UTF-8 text without retaining an encoded duplicate."""

    def __init__(
        self,
        *,
        max_upload_bytes: int,
        max_upload_chars: int,
        max_document_bytes: int,
        max_document_chars: int,
    ):
        self.max_upload_bytes = max_upload_bytes
        self.max_upload_chars = max_upload_chars
        self.max_document_bytes = max_document_bytes
        self.max_document_chars = max_document_chars
        self.upload_bytes = 0
        self.upload_chars = 0

    def add_text(self, text: Optional[str], label: str) -> Optional[str]:
        if text is None:
            return None
        if not isinstance(text, str):
            text = str(text)
        start, end = _trimmed_span(text)
        return self.add_span(text, start, end, label)

    def add_span(self, text: str, start: int, end: int, label: str) -> Optional[str]:
        if start >= end:
            return None

        char_count = end - start
        if char_count > self.max_document_chars:
            raise UploadValidationError(
                f"Extracted text for {label} exceeds the per-document character "
                f"limit ({self.max_document_chars:,})."
            )
        if self.upload_chars + char_count > self.max_upload_chars:
            raise UploadValidationError(
                f"Extracted text exceeds the per-upload character limit "
                f"({self.max_upload_chars:,})."
            )

        # Only make a substring after the character checks above.  Encoding is
        # bounded by max_document_chars and is released immediately after the
        # byte count, so it cannot create an unbounded second document copy.
        value = text[start:end]
        byte_count = len(value.encode("utf-8"))
        if byte_count > self.max_document_bytes:
            raise UploadValidationError(
                f"Extracted text for {label} exceeds the per-document byte "
                f"limit ({self.max_document_bytes:,})."
            )
        if self.upload_bytes + byte_count > self.max_upload_bytes:
            raise UploadValidationError(
                f"Extracted text exceeds the per-upload byte limit "
                f"({self.max_upload_bytes:,})."
            )

        self.upload_chars += char_count
        self.upload_bytes += byte_count
        return value


class DocumentUploader:
    """
    Handles document upload, parsing, chunking, and indexing into
    a temporary vector store for querying.
    """

    SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".md", ".csv", ".docx"}

    def __init__(
        self,
        embedding_model=None,
        max_bytes: int = MAX_UPLOAD_BYTES,
        *,
        max_extracted_text_bytes_per_upload: int = DEFAULT_MAX_EXTRACTED_TEXT_BYTES_PER_UPLOAD,
        max_extracted_text_chars_per_upload: int = DEFAULT_MAX_EXTRACTED_TEXT_CHARS_PER_UPLOAD,
        max_extracted_text_bytes_per_document: int = DEFAULT_MAX_EXTRACTED_TEXT_BYTES_PER_DOCUMENT,
        max_extracted_text_chars_per_document: int = DEFAULT_MAX_EXTRACTED_TEXT_CHARS_PER_DOCUMENT,
        max_pdf_pages: int = DEFAULT_MAX_PDF_PAGES,
        max_parsed_documents: int = DEFAULT_MAX_PARSED_DOCUMENTS,
        max_docx_entries: int = DEFAULT_MAX_DOCX_ENTRIES,
        max_docx_uncompressed_bytes: int = DEFAULT_MAX_DOCX_UNCOMPRESSED_BYTES,
        max_docx_member_bytes: int = DEFAULT_MAX_DOCX_MEMBER_BYTES,
        max_docx_compression_ratio: float = DEFAULT_MAX_DOCX_COMPRESSION_RATIO,
    ):
        self.embedding_model = embedding_model
        self.max_bytes = self._positive_limit("max_bytes", max_bytes)
        self.max_extracted_text_bytes_per_upload = self._positive_limit(
            "max_extracted_text_bytes_per_upload", max_extracted_text_bytes_per_upload
        )
        self.max_extracted_text_chars_per_upload = self._positive_limit(
            "max_extracted_text_chars_per_upload", max_extracted_text_chars_per_upload
        )
        self.max_extracted_text_bytes_per_document = self._positive_limit(
            "max_extracted_text_bytes_per_document", max_extracted_text_bytes_per_document
        )
        self.max_extracted_text_chars_per_document = self._positive_limit(
            "max_extracted_text_chars_per_document", max_extracted_text_chars_per_document
        )
        self.max_pdf_pages = self._positive_limit("max_pdf_pages", max_pdf_pages)
        self.max_parsed_documents = self._positive_limit(
            "max_parsed_documents", max_parsed_documents
        )
        self.max_docx_entries = self._positive_limit("max_docx_entries", max_docx_entries)
        self.max_docx_uncompressed_bytes = self._positive_limit(
            "max_docx_uncompressed_bytes", max_docx_uncompressed_bytes
        )
        self.max_docx_member_bytes = self._positive_limit(
            "max_docx_member_bytes", max_docx_member_bytes
        )
        self.max_docx_compression_ratio = self._positive_limit(
            "max_docx_compression_ratio", max_docx_compression_ratio
        )
        self._vectorstore: Optional[Chroma] = None
        self._owned_vectorstore: Optional[Chroma] = None
        self._chunks: List[Document] = []
        self._filename: str = ""

    @staticmethod
    def _positive_limit(name: str, value):
        try:
            valid = math.isfinite(value) and value > 0
        except (TypeError, ValueError):
            valid = False
        if not valid:
            raise ValueError(f"{name} must be positive")
        return value

    def _text_budget(self) -> _TextBudget:
        return _TextBudget(
            max_upload_bytes=self.max_extracted_text_bytes_per_upload,
            max_upload_chars=self.max_extracted_text_chars_per_upload,
            max_document_bytes=self.max_extracted_text_bytes_per_document,
            max_document_chars=self.max_extracted_text_chars_per_document,
        )

    def _append_document(
        self,
        docs: List[Document],
        text: Optional[str],
        metadata: dict,
        budget: _TextBudget,
        label: str,
        *,
        start: Optional[int] = None,
        end: Optional[int] = None,
    ) -> None:
        if start is not None and end is not None:
            value = budget.add_span(text or "", start, end, label)
        else:
            value = budget.add_text(text, label)
        if value is None:
            return
        if len(docs) >= self.max_parsed_documents:
            raise UploadValidationError(
                f"Parsed document count exceeds the per-upload limit "
                f"({self.max_parsed_documents:,})."
            )
        docs.append(Document(page_content=value, metadata=metadata))

    @property
    def is_ready(self) -> bool:
        return self._vectorstore is not None and len(self._chunks) > 0

    @property
    def filename(self) -> str:
        return self._filename

    @property
    def num_chunks(self) -> int:
        return len(self._chunks)

    @property
    def vectorstore(self) -> Optional[Chroma]:
        """The in-memory vector store for the active document (None until indexed)."""
        return self._vectorstore

    @property
    def chunks(self) -> List[Document]:
        """Return the active indexed chunks for architectures needing a corpus."""
        return list(self._chunks)

    # ── Validation ──────────────────────────────────────────────────────────
    def _validate(self, file_content: bytes, filename: str) -> str:
        if not file_content:
            raise UploadValidationError("Empty upload (0 bytes).")
        if len(file_content) > self.max_bytes:
            raise UploadValidationError(
                f"File too large: {len(file_content):,} bytes "
                f"(max allowed: {self.max_bytes:,} bytes)."
            )
        safe_name = _safe_basename(filename)
        ext = os.path.splitext(safe_name)[1].lower()
        if ext not in self.SUPPORTED_EXTENSIONS:
            raise UploadValidationError(
                f"Unsupported file type '{ext}'. Supported formats: {sorted(self.SUPPORTED_EXTENSIONS)}"
            )

        # Magic-byte sniff for binary types within first 1KB (tolerates BOM/leading whitespace)
        header_sample = file_content[:1024]
        if ext == ".pdf" and b"%PDF" not in header_sample:
            raise UploadValidationError("File is not a valid PDF (missing %PDF signature header).")
        if ext == ".docx" and b"PK\x03\x04" not in header_sample:
            raise UploadValidationError("File is not a valid DOCX (missing ZIP/DOCX header).")
        return ext

    # ── Parsers ─────────────────────────────────────────────────────────────
    def parse_file(self, file_content: bytes, filename: str) -> List[Document]:
        ext = self._validate(file_content, filename)
        previous_filename = self._filename
        self._filename = _safe_basename(filename)

        try:
            if ext == ".pdf":
                return self._parse_pdf(file_content)
            if ext == ".docx":
                return self._parse_docx(file_content)
            return self._parse_text(file_content)
        except Exception:
            # Direct parse_file callers should get the same transactional
            # filename behavior as index_documents callers using a staged copy.
            self._filename = previous_filename
            raise

    @staticmethod
    def _decode_text(content: bytes) -> str:
        """Decode once in the common case, with bounded fallback attempts."""
        if content.startswith((b"\xff\xfe", b"\xfe\xff")):
            encodings = ("utf-16", "utf-8-sig", "cp1252", "latin-1")
        elif content.startswith(b"\xef\xbb\xbf"):
            encodings = ("utf-8-sig", "utf-8", "cp1252", "latin-1")
        else:
            encodings = ("utf-8", "cp1252", "latin-1", "utf-16")

        for encoding in encodings:
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="replace")

    def _parse_text(self, content: bytes) -> List[Document]:
        # A text upload is already its own extracted representation.  Reject it
        # before decoding when the raw bytes exceed the extraction budget; this
        # prevents a large invalid-encoding payload from being decoded repeatedly.
        if len(content) > self.max_extracted_text_bytes_per_upload:
            raise UploadValidationError(
                "Extracted text exceeds the per-upload byte limit "
                f"({self.max_extracted_text_bytes_per_upload:,})."
            )

        text = self._decode_text(content)
        if len(text) > self.max_extracted_text_chars_per_upload:
            raise UploadValidationError(
                "Extracted text exceeds the per-upload character limit "
                f"({self.max_extracted_text_chars_per_upload:,})."
            )

        budget = self._text_budget()
        docs: List[Document] = []
        cursor = 0
        paragraph_index = 0
        while cursor <= len(text):
            separator = text.find("\n\n", cursor)
            paragraph_end = len(text) if separator < 0 else separator
            start, end = _trimmed_span(text, cursor, paragraph_end)
            # Preserve the existing behavior: normal paragraphs shorter than
            # eleven characters are ignored, with a whole-file fallback below.
            if end - start > 10:
                self._append_document(
                    docs,
                    text,
                    {"source": self._filename, "paragraph": paragraph_index},
                    budget,
                    f"text paragraph {paragraph_index}",
                    start=start,
                    end=end,
                )
            paragraph_index += 1
            if separator < 0:
                break
            cursor = separator + 2

        if not docs:
            start, end = _trimmed_span(text)
            if start < end:
                self._append_document(
                    docs,
                    text,
                    {"source": self._filename, "paragraph": 0},
                    budget,
                    "text upload",
                    start=start,
                    end=end,
                )
        if not docs:
            raise UploadValidationError("File contains no readable text.")
        return docs

    @staticmethod
    def _pdf_page_count(pdf_doc) -> int:
        page_count = getattr(pdf_doc, "page_count", None)
        return int(page_count) if page_count is not None else len(pdf_doc)

    def _parse_pdf(self, content: bytes) -> List[Document]:
        docs: List[Document] = []
        parse_errors: List[str] = []
        budget = self._text_budget()

        # 1. Primary engine: PyMuPDF (fast, handles complex layouts, fonts, encodings)
        pdf_doc = None
        try:
            import pymupdf

            pdf_doc = pymupdf.open(stream=content, filetype="pdf")
            if pdf_doc.is_encrypted:
                try:
                    pdf_doc.authenticate("")
                except Exception:
                    pass

            page_count = self._pdf_page_count(pdf_doc)
            if page_count > self.max_pdf_pages:
                raise UploadValidationError(
                    f"PDF page count {page_count:,} exceeds the per-upload limit "
                    f"({self.max_pdf_pages:,})."
                )

            for i, page in enumerate(pdf_doc):
                try:
                    self._append_document(
                        docs,
                        page.get_text("text"),
                        {"source": self._filename, "page": i + 1},
                        budget,
                        f"PDF page {i + 1}",
                    )
                except UploadValidationError:
                    raise
                except Exception as page_exc:
                    logger.warning("PyMuPDF extraction failed on page %d: %s", i + 1, page_exc)
        except UploadValidationError:
            raise
        except Exception as exc:
            diagnostic = str(exc)[:500]
            parse_errors.append(f"PyMuPDF error: {diagnostic}")
            logger.warning("PyMuPDF parser failed, trying pypdf fallback: %s", exc)
        finally:
            if pdf_doc is not None:
                try:
                    pdf_doc.close()
                except Exception:
                    logger.warning("Could not close PyMuPDF document", exc_info=True)

        # 2. Secondary engine fallback: pypdf
        if not docs:
            pdf_stream = io.BytesIO(content)
            try:
                from pypdf import PdfReader
                reader = PdfReader(pdf_stream)
                if reader.is_encrypted:
                    try:
                        reader.decrypt("")
                    except Exception:
                        pass

                pages = reader.pages
                page_count = len(pages)
                if page_count > self.max_pdf_pages:
                    raise UploadValidationError(
                        f"PDF page count {page_count:,} exceeds the per-upload limit "
                        f"({self.max_pdf_pages:,})."
                    )

                for i, page in enumerate(pages):
                    try:
                        self._append_document(
                            docs,
                            page.extract_text() or "",
                            {"source": self._filename, "page": i + 1},
                            budget,
                            f"PDF page {i + 1}",
                        )
                    except UploadValidationError:
                        raise
                    except Exception as page_exc:
                        logger.warning("pypdf page %d extraction failed: %s", i + 1, page_exc)
            except UploadValidationError:
                raise
            except Exception as exc:
                diagnostic = str(exc)[:500]
                parse_errors.append(f"pypdf error: {diagnostic}")
                logger.warning("pypdf fallback failed: %s", exc)
            finally:
                pdf_stream.close()

        if not docs:
            err_msg = (
                "PDF contained no extractable text layer. "
                "This typically happens if the PDF consists of scanned images/photos without OCR. "
                "Please upload a text-searchable PDF or a document in .txt, .md, or .docx format."
            )
            if parse_errors:
                err_msg += f" (Diagnostics: {'; '.join(parse_errors)})"
            raise UploadValidationError(err_msg)

        return docs

    def _inspect_docx_archive(self, content: bytes) -> None:
        """Validate DOCX ZIP metadata without reading or extracting members."""
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                infos = archive.infolist()
                if len(infos) > self.max_docx_entries:
                    raise UploadValidationError(
                        f"DOCX contains too many ZIP entries ({len(infos):,}; "
                        f"max {self.max_docx_entries:,})."
                    )

                total_uncompressed = 0
                for info in infos:
                    name = info.filename
                    normalized = name.replace("\\", "/")
                    posix_name = PurePosixPath(normalized)
                    windows_name = PureWindowsPath(normalized)
                    if (
                        not name
                        or "\x00" in name
                        or normalized.startswith("/")
                        or posix_name.is_absolute()
                        or windows_name.drive
                        or windows_name.root
                        or any(part == ".." for part in posix_name.parts)
                    ):
                        raise UploadValidationError(
                            f"DOCX ZIP member has an unsafe path: {name!r}."
                        )

                    if info.file_size < 0 or info.compress_size < 0:
                        raise UploadValidationError("DOCX ZIP contains invalid size metadata.")
                    if info.file_size > self.max_docx_member_bytes:
                        raise UploadValidationError(
                            f"DOCX ZIP member {name!r} is too large "
                            f"({info.file_size:,} bytes; max "
                            f"{self.max_docx_member_bytes:,})."
                        )

                    total_uncompressed += info.file_size
                    if total_uncompressed > self.max_docx_uncompressed_bytes:
                        raise UploadValidationError(
                            "DOCX ZIP expands beyond the total uncompressed-size "
                            f"limit ({self.max_docx_uncompressed_bytes:,} bytes)."
                        )

                    if info.file_size:
                        if info.compress_size == 0:
                            raise UploadValidationError(
                                f"DOCX ZIP member {name!r} has an invalid compression size."
                            )
                        ratio = info.file_size / info.compress_size
                        if ratio > self.max_docx_compression_ratio:
                            raise UploadValidationError(
                                f"DOCX ZIP member {name!r} has a suspicious compression "
                                f"ratio ({ratio:,.1f}:1; max "
                                f"{self.max_docx_compression_ratio:,.1f}:1)."
                            )
        except UploadValidationError:
            raise
        except Exception as exc:
            diagnostic = str(exc)[:500]
            raise UploadValidationError(f"Could not inspect DOCX ZIP: {diagnostic}") from exc

    def _table_row_text(self, row) -> str:
        """Join a table row only after checking its bounded character size."""
        parts: List[str] = []
        char_count = 0
        for cell in row.cells:
            cell_text = cell.text
            if not isinstance(cell_text, str):
                cell_text = str(cell_text)
            start, end = _trimmed_span(cell_text)
            if start >= end:
                continue
            separator_chars = 3 if parts else 0  # `` | ``
            cell_chars = end - start
            char_count += separator_chars + cell_chars
            if char_count > self.max_extracted_text_chars_per_document:
                raise UploadValidationError(
                    "DOCX table row exceeds the per-document character limit "
                    f"({self.max_extracted_text_chars_per_document:,})."
                )
            parts.append(cell_text[start:end])
        return " | ".join(parts)

    def _parse_docx(self, content: bytes) -> List[Document]:
        self._inspect_docx_archive(content)

        from docx import Document as DocxDocument

        archive_stream = io.BytesIO(content)
        try:
            try:
                doc = DocxDocument(archive_stream)
            except Exception as exc:
                diagnostic = str(exc)[:500]
                raise UploadValidationError(f"Could not parse DOCX: {diagnostic}") from exc

            budget = self._text_budget()
            docs: List[Document] = []
            paragraph_index = 0
            for para in doc.paragraphs:
                text = para.text
                if isinstance(text, str):
                    start, end = _trimmed_span(text)
                    if end - start > 10:
                        self._append_document(
                            docs,
                            text,
                            {"source": self._filename, "paragraph": paragraph_index},
                            budget,
                            f"DOCX paragraph {paragraph_index}",
                            start=start,
                            end=end,
                        )
                paragraph_index += 1

            for table in doc.tables:
                for row in table.rows:
                    row_text = self._table_row_text(row)
                    if len(row_text) > 10:
                        self._append_document(
                            docs,
                            row_text,
                            {"source": self._filename, "paragraph": paragraph_index},
                            budget,
                            f"DOCX table row {paragraph_index}",
                        )
                    paragraph_index += 1
        except UploadValidationError:
            raise
        except Exception as exc:
            diagnostic = str(exc)[:500]
            raise UploadValidationError(f"Could not parse DOCX: {diagnostic}") from exc
        finally:
            archive_stream.close()

        if not docs:
            raise UploadValidationError("DOCX document contained no text content.")
        return docs

    # ── Indexing ────────────────────────────────────────────────────────────
    def index_documents(
        self,
        file_content: bytes,
        filename: str,
        chunk_size: int = CHUNK_SIZE,
        chunk_overlap: int = CHUNK_OVERLAP,
    ) -> int:
        # Parsers use _filename for metadata. Stage them on a copy so a failed
        # parse, split, or embedding cannot change the active upload's state.
        staged_parser = copy(self)
        docs = staged_parser.parse_file(file_content, filename)

        splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            separators=["\n\n", "\n", ". ", " ", ""],
        )
        chunks = splitter.split_documents(docs)

        embedding_model = self.embedding_model
        if embedding_model is None:
            from module1.data_preprocessing import get_embedding_model
            embedding_model = get_embedding_model()

        # Ephemeral Chroma clients share collections by name. Give each upload
        # its own UUID and retain a handle to clean up even if embedding fails.
        vectorstore = Chroma(
            collection_name=f"upload-{uuid4().hex}",
            embedding_function=embedding_model,
        )
        try:
            vectorstore.add_documents(chunks)
        except Exception:
            self._delete_owned_collection(vectorstore)
            raise

        previous_owned_store = self._owned_vectorstore
        self._vectorstore = vectorstore
        self._owned_vectorstore = vectorstore
        self._chunks = chunks
        self._filename = staged_parser.filename
        self.embedding_model = embedding_model
        self._delete_owned_collection(previous_owned_store)
        logger.info("Indexed %d chunks from %s", len(self._chunks), self._filename)
        return len(self._chunks)

    @staticmethod
    def _delete_owned_collection(vectorstore: Optional[Chroma]) -> None:
        """Only call for UUID-named stores created by this uploader."""
        if vectorstore is not None:
            try:
                vectorstore.delete_collection()
            except Exception:
                # Cleanup must not mask an indexing failure or turn a completed
                # replacement into a reported failure after it was published.
                logger.warning("Could not delete an owned upload collection", exc_info=True)

    # ── Retrieval ───────────────────────────────────────────────────────────
    def search(self, query: str, k: int = 5) -> List[Document]:
        if not self.is_ready:
            return []
        return self._vectorstore.similarity_search(query, k=k)

    def search_with_scores(self, query: str, k: int = 5) -> List[tuple]:
        if not self.is_ready:
            return []
        return self._vectorstore.similarity_search_with_score(query, k=k)

    def clear(self) -> None:
        owned_store = self._owned_vectorstore
        self._vectorstore = None
        self._owned_vectorstore = None
        self._chunks = []
        self._filename = ""
        self._delete_owned_collection(owned_store)
