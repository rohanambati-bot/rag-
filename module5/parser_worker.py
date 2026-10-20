"""Lightweight parser and executable worker; no application/ML imports.

Run this file directly, never ``python -m module5.parser_worker``: module5's
package initializer imports the application stack. The executable bootstrap sets
resource limits before reading input or importing PDF/DOCX/ZIP libraries.
"""

from __future__ import annotations

import io
import json
import os
import sys
from dataclasses import dataclass
from pathlib import PurePosixPath, PureWindowsPath
from typing import Optional

DEFAULT_MAX_EXTRACTED_TEXT_BYTES_PER_UPLOAD = 8 * 1024 * 1024
DEFAULT_MAX_EXTRACTED_TEXT_CHARS_PER_UPLOAD = 2_000_000
DEFAULT_MAX_EXTRACTED_TEXT_BYTES_PER_DOCUMENT = 2 * 1024 * 1024
DEFAULT_MAX_EXTRACTED_TEXT_CHARS_PER_DOCUMENT = 500_000
DEFAULT_MAX_PDF_PAGES = 200
DEFAULT_MAX_PARSED_DOCUMENTS = 10_000
DEFAULT_MAX_DOCX_ENTRIES = 1_024
DEFAULT_MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
DEFAULT_MAX_DOCX_MEMBER_BYTES = 10 * 1024 * 1024
DEFAULT_MAX_DOCX_COMPRESSION_RATIO = 100.0

POLICY_DEFAULTS = {
    "max_bytes": 50 * 1024 * 1024,
    "max_extracted_text_bytes_per_upload": DEFAULT_MAX_EXTRACTED_TEXT_BYTES_PER_UPLOAD,
    "max_extracted_text_chars_per_upload": DEFAULT_MAX_EXTRACTED_TEXT_CHARS_PER_UPLOAD,
    "max_extracted_text_bytes_per_document": DEFAULT_MAX_EXTRACTED_TEXT_BYTES_PER_DOCUMENT,
    "max_extracted_text_chars_per_document": DEFAULT_MAX_EXTRACTED_TEXT_CHARS_PER_DOCUMENT,
    "max_pdf_pages": DEFAULT_MAX_PDF_PAGES,
    "max_parsed_documents": DEFAULT_MAX_PARSED_DOCUMENTS,
    "max_docx_entries": DEFAULT_MAX_DOCX_ENTRIES,
    "max_docx_uncompressed_bytes": DEFAULT_MAX_DOCX_UNCOMPRESSED_BYTES,
    "max_docx_member_bytes": DEFAULT_MAX_DOCX_MEMBER_BYTES,
    "max_docx_compression_ratio": DEFAULT_MAX_DOCX_COMPRESSION_RATIO,
}

# Only these fixed messages cross the process boundary. Native exceptions, paths,
# archive names, parser diagnostics and uploaded contents never become errors.
ERROR_MESSAGES = {
    "upload_bytes": "File too large for the encoded upload limit.",
    "upload_text_bytes": "Extracted text exceeds the per-upload byte limit.",
    "upload_text_chars": "Extracted text exceeds the per-upload character limit.",
    "document_bytes": "Extracted text exceeds the per-document byte limit.",
    "document_chars": "Extracted text exceeds the per-document character limit.",
    "documents": "Parsed document count exceeds the per-upload limit.",
    "pages": "PDF page count exceeds the per-upload limit.",
    "pdf": "PDF contains no readable text or could not be parsed safely.",
    "text": "File contains no readable text.",
    "entries": "DOCX contains too many ZIP entries.",
    "path": "DOCX ZIP member has an unsafe path.",
    "member": "DOCX ZIP member is too large.",
    "expanded": "DOCX ZIP exceeds the total uncompressed-size limit.",
    "ratio": "DOCX ZIP member exceeds the compression ratio limit.",
    "zip": "Could not inspect DOCX ZIP safely.",
    "docx": "DOCX could not be parsed safely or contains no readable text.",
    "invalid": "Document could not be parsed safely.",
}


class UploadValidationError(ValueError):
    """A public, bounded upload rejection."""


def _reject(code: str) -> None:
    raise UploadValidationError(ERROR_MESSAGES[code])


def _safe_basename(name: str) -> str:
    base = PurePosixPath(name.replace("\\", "/")).name
    base = "".join(c for c in base if 32 <= ord(c) < 127)
    return base[:255] if base and base not in {".", ".."} else "uploaded"


def _trimmed_span(text: str, start: int = 0, end: Optional[int] = None) -> tuple[int, int]:
    if end is None:
        end = len(text)
    while start < end and text[start].isspace():
        start += 1
    while end > start and text[end - 1].isspace():
        end -= 1
    return start, end


@dataclass
class ParsedDocument:
    page_content: str
    metadata: dict


class _TextBudget:
    def __init__(self, *, max_upload_bytes, max_upload_chars, max_document_bytes, max_document_chars):
        self.max_upload_bytes = max_upload_bytes
        self.max_upload_chars = max_upload_chars
        self.max_document_bytes = max_document_bytes
        self.max_document_chars = max_document_chars
        self.upload_bytes = 0
        self.upload_chars = 0

    def add_text(self, text, label):
        if text is None:
            return None
        if not isinstance(text, str):
            text = str(text)
        return self.add_span(text, *_trimmed_span(text), label)

    def add_span(self, text, start, end, label):
        if start >= end:
            return None
        count = end - start
        if count > self.max_document_chars:
            _reject("document_chars")
        if self.upload_chars + count > self.max_upload_chars:
            _reject("upload_text_chars")
        value = text[start:end]
        byte_count = len(value.encode("utf-8"))
        if byte_count > self.max_document_bytes:
            _reject("document_bytes")
        if self.upload_bytes + byte_count > self.max_upload_bytes:
            _reject("upload_text_bytes")
        self.upload_bytes += byte_count
        self.upload_chars += count
        return value


class BoundedParser:
    """Parser internals, shared with helper tests; production calls the worker."""

    def __init__(self, filename: str, policy: dict):
        if policy.keys() != POLICY_DEFAULTS.keys():
            _reject("invalid")
        for name, value in policy.items():
            import math
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value <= 0:
                _reject("invalid")
            if name != "max_docx_compression_ratio" and not isinstance(value, int):
                _reject("invalid")
            setattr(self, name, value)
        self._filename = _safe_basename(filename)

    def _text_budget(self):
        return _TextBudget(
            max_upload_bytes=self.max_extracted_text_bytes_per_upload,
            max_upload_chars=self.max_extracted_text_chars_per_upload,
            max_document_bytes=self.max_extracted_text_bytes_per_document,
            max_document_chars=self.max_extracted_text_chars_per_document,
        )

    def _append_document(self, docs, text, metadata, budget, label, *, start=None, end=None):
        value = budget.add_span(text or "", start, end, label) if start is not None and end is not None else budget.add_text(text, label)
        if value is None:
            return
        if len(docs) >= self.max_parsed_documents:
            _reject("documents")
        docs.append(ParsedDocument(value, metadata))

    @staticmethod
    def _decode_text(content):
        if content.startswith((b"\xff\xfe", b"\xfe\xff")):
            encodings = ("utf-16", "utf-8-sig", "cp1252", "latin-1")
        elif content.startswith(b"\xef\xbb\xbf"):
            encodings = ("utf-8-sig", "utf-8", "cp1252", "latin-1")
        else:
            encodings = ("utf-8", "cp1252", "latin-1", "utf-16")
        for encoding in encodings:
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="replace")

    def _parse_text(self, content):
        if len(content) > self.max_extracted_text_bytes_per_upload:
            _reject("upload_text_bytes")
        text = self._decode_text(content)
        if len(text) > self.max_extracted_text_chars_per_upload:
            _reject("upload_text_chars")
        budget, docs, cursor, paragraph = self._text_budget(), [], 0, 0
        while cursor <= len(text):
            separator = text.find("\n\n", cursor)
            start, end = _trimmed_span(text, cursor, len(text) if separator < 0 else separator)
            if end - start > 10:
                self._append_document(docs, text, {"source": self._filename, "paragraph": paragraph}, budget, "text", start=start, end=end)
            paragraph += 1
            if separator < 0:
                break
            cursor = separator + 2
        if not docs:
            self._append_document(docs, text, {"source": self._filename, "paragraph": 0}, budget, "text")
        if not docs:
            _reject("text")
        return docs

    @staticmethod
    def _pdf_page_count(pdf_doc):
        count = getattr(pdf_doc, "page_count", None)
        return int(count) if count is not None else len(pdf_doc)

    def _parse_pdf(self, content):
        docs, budget, pdf_doc = [], self._text_budget(), None
        try:
            import pymupdf
            pdf_doc = pymupdf.open(stream=content, filetype="pdf")
            if pdf_doc.is_encrypted:
                pdf_doc.authenticate("")
            if self._pdf_page_count(pdf_doc) > self.max_pdf_pages:
                _reject("pages")
            for index, page in enumerate(pdf_doc):
                try:
                    self._append_document(docs, page.get_text("text"), {"source": self._filename, "page": index + 1}, budget, "PDF page")
                except UploadValidationError:
                    raise
                except Exception:
                    continue
        except UploadValidationError:
            raise
        except Exception:
            pass
        finally:
            if pdf_doc is not None:
                try:
                    pdf_doc.close()
                except Exception:
                    pass
        if not docs:
            with io.BytesIO(content) as stream:
                try:
                    from pypdf import PdfReader
                    reader = PdfReader(stream)
                    if reader.is_encrypted:
                        reader.decrypt("")
                    if len(reader.pages) > self.max_pdf_pages:
                        _reject("pages")
                    for index, page in enumerate(reader.pages):
                        try:
                            self._append_document(docs, page.extract_text() or "", {"source": self._filename, "page": index + 1}, budget, "PDF page")
                        except UploadValidationError:
                            raise
                        except Exception:
                            continue
                except UploadValidationError:
                    raise
                except Exception:
                    pass
        if not docs:
            _reject("pdf")
        return docs

    def _inspect_docx_archive(self, content):
        # No extraction to disk. Inspect every entry, then stream/CRC-check its
        # actual decompressed bytes before python-docx opens the same container.
        import zipfile
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                infos = archive.infolist()
                if len(infos) > self.max_docx_entries:
                    _reject("entries")
                total, names = 0, set()
                for info in infos:
                    name = info.filename
                    normalized = name.replace("\\", "/")
                    path, winpath = PurePosixPath(normalized), PureWindowsPath(normalized)
                    if not name or "\x00" in info.orig_filename or normalized.startswith("/") or path.is_absolute() or winpath.drive or winpath.root or ".." in path.parts:
                        _reject("path")
                    if normalized in names or info.flag_bits & 1 or (info.external_attr >> 16) & 0o170000 == 0o120000:
                        _reject("zip")
                    names.add(normalized)
                    if info.file_size < 0 or info.compress_size < 0 or info.compress_size > len(content):
                        _reject("zip")
                    if info.file_size > self.max_docx_member_bytes:
                        _reject("member")
                    total += info.file_size
                    if total > self.max_docx_uncompressed_bytes:
                        _reject("expanded")
                    if info.file_size and (not info.compress_size or info.file_size / info.compress_size > self.max_docx_compression_ratio):
                        _reject("ratio")
                actual_total = 0
                for info in infos:
                    actual = 0
                    with archive.open(info) as member:
                        while True:
                            block = member.read(64 * 1024)
                            if not block:
                                break
                            actual += len(block)
                            actual_total += len(block)
                            if actual > self.max_docx_member_bytes:
                                _reject("member")
                            if actual_total > self.max_docx_uncompressed_bytes:
                                _reject("expanded")
                    if actual != info.file_size:
                        _reject("zip")
        except UploadValidationError:
            raise
        except Exception:
            _reject("zip")

    def _table_row_text(self, row):
        parts, count = [], 0
        for cell in row.cells:
            text = str(cell.text)
            start, end = _trimmed_span(text)
            if start >= end:
                continue
            count += end - start + (3 if parts else 0)
            if count > self.max_extracted_text_chars_per_document:
                _reject("document_chars")
            parts.append(text[start:end])
        return " | ".join(parts)

    def _parse_docx(self, content):
        self._inspect_docx_archive(content)
        from docx import Document as DocxDocument
        try:
            with io.BytesIO(content) as stream:
                document = DocxDocument(stream)
                docs, budget, index = [], self._text_budget(), 0
                for paragraph in document.paragraphs:
                    text = paragraph.text
                    start, end = _trimmed_span(text)
                    if end - start > 10:
                        self._append_document(docs, text, {"source": self._filename, "paragraph": index}, budget, "DOCX paragraph", start=start, end=end)
                    index += 1
                for table in document.tables:
                    for row in table.rows:
                        text = self._table_row_text(row)
                        if len(text) > 10:
                            self._append_document(docs, text, {"source": self._filename, "paragraph": index}, budget, "DOCX row")
                        index += 1
        except UploadValidationError:
            raise
        except Exception:
            _reject("docx")
        if not docs:
            _reject("docx")
        return docs


def _apply_resource_limits(cpu_seconds: int, memory_bytes: int, output_bytes: int) -> dict:
    import resource
    if os.name != "posix" or sys.platform not in {"linux", "darwin"}:
        raise RuntimeError("unsupported platform")
    requested = {
        "RLIMIT_CPU": (cpu_seconds, cpu_seconds),
        "RLIMIT_CORE": (0, 0),
        "RLIMIT_FSIZE": (output_bytes, output_bytes),
        "RLIMIT_NOFILE": (64, 64),
    }
    # Linux enforces address-space limits. macOS does not reliably enforce AS or
    # RSS, and DATA is capped by a low kernel maximum on some releases. The
    # supervising parent establishes libproc RSS monitoring before starting us.
    if sys.platform == "linux":
        requested["RLIMIT_AS"] = (memory_bytes, memory_bytes)
    applied = {}
    for name, (soft, hard) in requested.items():
        key = getattr(resource, name)
        _, current_hard = resource.getrlimit(key)
        if current_hard != resource.RLIM_INFINITY:
            hard = min(hard, current_hard)
            soft = min(soft, hard)
        resource.setrlimit(key, (soft, hard))
        applied[name] = list(resource.getrlimit(key))
    return applied


def main() -> int:
    try:
        # This must precede any native parser imports or untrusted input reads.
        cpu, memory, output_limit = (int(value) for value in sys.argv[1:])
        if min(cpu, memory, output_limit) <= 0:
            return 1
        limits = _apply_resource_limits(cpu, memory, output_limit)
        request_raw = sys.stdin.buffer.read(16 * 1024 + 1)
        if len(request_raw) > 16 * 1024:
            return 1
        request = json.loads(request_raw)
        parser = BoundedParser(request["filename"], request["policy"])
        with open("upload.bin", "rb") as upload:
            content = upload.read(parser.max_bytes + 1)
        if not content or len(content) > parser.max_bytes:
            _reject("upload_bytes")
        extension = request["extension"]
        if extension == ".pdf":
            docs = parser._parse_pdf(content)
        elif extension == ".docx":
            docs = parser._parse_docx(content)
        elif extension in {".txt", ".md", ".csv"}:
            docs = parser._parse_text(content)
        else:
            _reject("invalid")
        response = {
            "version": 1,
            "documents": [{"page_content": doc.page_content, "metadata": doc.metadata} for doc in docs],
            "worker": {
                "pid": os.getpid(), "parent_pid": os.getppid(), "limits": limits,
                "heavy_stack_loaded": any(name in sys.modules for name in ("torch", "chromadb", "langchain_community")),
            },
        }
    except UploadValidationError as exc:
        code = next((key for key, message in ERROR_MESSAGES.items() if message == str(exc)), "invalid")
        response = {"version": 1, "error": code}
    except BaseException:
        response = {"version": 1, "error": "invalid"}
    try:
        # Streaming encoding avoids constructing an unbounded encoded duplicate.
        written = 0
        for fragment in json.JSONEncoder(ensure_ascii=False, separators=(",", ":")).iterencode(response):
            encoded = fragment.encode("utf-8")
            written += len(encoded)
            if written > output_limit:
                return 1
            sys.stdout.buffer.write(encoded)
        sys.stdout.buffer.flush()
        return 0
    except BaseException:
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
